"""Prueba E2E en navegador real (Edge vía Playwright) contra un servidor local ya en marcha.

Uso (con el backend corriendo en :8000 y una cuenta admin):
  set E2E_ADMIN_USER=demo & set E2E_ADMIN_PASS=... & python tests/e2e/e2e_visual.py OUT_DIR
Genera capturas y verifica comportamiento. No forma parte de pytest (necesita navegador y servidor).
"""
import io
import os
import sys
import time

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

BASE = os.environ.get("E2E_BASE", "http://127.0.0.1:8000")
OUT = sys.argv[1] if len(sys.argv) > 1 else "e2e_out"
USER = os.environ["E2E_ADMIN_USER"]
PASS = os.environ["E2E_ADMIN_PASS"]
os.makedirs(OUT, exist_ok=True)
results = []


def check(name, cond, extra=""):
    results.append((name, bool(cond)))
    print(("PASS " if cond else "FAIL ") + name + (" · " + str(extra) if extra else ""))


def fill_form(page, name="Laura Hernández Ruiz", phone="449 555 0182", email="laura@example.com", n=3):
    page.fill("#full_name", name)
    page.fill("#phone", phone)
    page.fill("#email", email)
    page.fill("#attendees", str(n))
    page.select_option("#arrival_window", "12-15")
    page.check("#privacy_accepted")


with sync_playwright() as p:
    browser = p.chromium.launch(channel="msedge", headless=True)

    # ------------------------------------------------ capturas responsivas
    for label, vp in {"desktop": (1440, 900), "tablet": (820, 1180), "mobile": (390, 844)}.items():
        ctx = browser.new_context(viewport={"width": vp[0], "height": vp[1]}, reduced_motion="reduce", device_scale_factor=1)
        page = ctx.new_page()
        errors = []
        page.on("console", lambda m: errors.append(m.text) if m.type in ("error", "warning") else None)
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.goto(BASE + "/?utm_source=facebook&utm_medium=paid&utm_campaign=expo", wait_until="networkidle")
        page.screenshot(path=os.path.join(OUT, "%s-top.png" % label))
        page.screenshot(path=os.path.join(OUT, "%s-full.png" % label), full_page=True)
        check("[%s] sin errores de consola/CSP" % label, not errors, errors[:3])
        sw = page.evaluate("document.documentElement.scrollWidth")
        check("[%s] sin desbordamiento horizontal" % label, sw <= vp[0] + 1, sw)
        ctx.close()

    # ------------------------------------------------ formulario (escritorio)
    ctx = browser.new_context(viewport={"width": 1440, "height": 900})
    page = ctx.new_page()
    posts = []
    page.on("request", lambda r: posts.append(r.url) if r.method == "POST" and r.url.endswith("/api/registrations") else None)
    page.goto(BASE + "/?utm_source=facebook&utm_medium=paid&utm_campaign=expo", wait_until="networkidle")

    page.click("a.btn[href='#registro']")  # CTA del hero hace scroll suave al formulario
    page.wait_for_timeout(1200)
    check("CTA desplaza hasta el formulario", page.evaluate("document.getElementById('registro').getBoundingClientRect().top") < 200)

    time.sleep(2.5)  # tiempo mínimo del formulario (antibot)
    page.click("#submit-btn")
    page.wait_for_timeout(300)
    check("validación: nombre vacío", "nombre" in page.inner_text("#err-full_name").lower())
    check("validación: teléfono vacío", "celular" in page.inner_text("#err-phone").lower())
    check("validación: correo vacío", "correo" in page.inner_text("#err-email").lower())
    check("validación: consentimiento", "Aviso de Privacidad" in page.inner_text("#err-privacy_accepted"))
    check("sin envío al servidor con datos inválidos", len(posts) == 0)
    page.locator("#panel").screenshot(path=os.path.join(OUT, "form-errors.png"))

    page.fill("#phone", "123")
    page.fill("#email", "no-es-correo")
    page.click("#submit-btn")
    check("validación: teléfono inválido", "10 dígitos" in page.inner_text("#err-phone"))
    check("validación: correo inválido", "válido" in page.inner_text("#err-email"))

    fill_form(page)
    # Stepper accesible
    page.fill("#attendees", "1")
    page.click("#att-plus"); page.click("#att-plus")
    check("stepper + suma", page.input_value("#attendees") == "3")
    page.click("#att-minus")
    check("stepper − resta", page.input_value("#attendees") == "2")
    page.click("#att-plus")
    page.locator("#panel").screenshot(path=os.path.join(OUT, "form-filled.png"))

    # doble clic rápido: debe producir UN solo POST
    page.dblclick("#submit-btn")
    page.wait_for_selector("#confirm:not([hidden])", timeout=8000)
    check("doble clic genera un único envío", len(posts) == 1, len(posts))
    folio = page.inner_text("#c-folio")
    check("confirmación con folio", folio.startswith("LANDA-"), folio)
    check("confirmación: total de asistentes", "3 personas" in page.inner_text("#c-attendees"))
    check("URL sin datos personales", "laura" not in page.url.lower() and "449" not in page.url)
    check("enlace de WhatsApp con el número correcto", "wa.me/524493481867" in page.get_attribute("#c-wa", "href"))
    check("enlace de Google Maps", "google.com/maps" in page.get_attribute("#c-maps", "href"))
    page.locator("#panel").screenshot(path=os.path.join(OUT, "confirmation.png"))
    with page.expect_download() as dl:
        page.click("#add-calendar")
    ics = open(dl.value.path(), encoding="utf-8").read()
    check("calendario .ics válido", "DTSTART:20261107T160000Z" in ics and folio in ics)
    ctx.grant_permissions(["clipboard-read", "clipboard-write"])
    page.click("#copy-folio")
    page.wait_for_timeout(300)
    check("copiar folio", page.evaluate("navigator.clipboard.readText()") == folio)
    ctx.close()

    # ------------------------------------------------ error temporal del servidor: no confirma y permite reintentar
    ctx = browser.new_context(viewport={"width": 390, "height": 844})
    page = ctx.new_page()
    page.goto(BASE + "/", wait_until="networkidle")
    time.sleep(2.5)
    fail_once = {"n": 0}

    def flaky(route):
        if fail_once["n"] == 0:
            fail_once["n"] += 1
            route.fulfill(status=503, content_type="application/json", body='{"detail":"x"}')
        else:
            route.continue_()

    page.route("**/api/registrations", flaky)
    fill_form(page, name="Pedro Gómez Salas", phone="+52 449 777 1010", email="pedro@example.com", n=2)
    page.click("#submit-btn")
    page.wait_for_timeout(800)
    check("error 503: NO se muestra confirmación", page.locator("#confirm").is_hidden())
    check("error 503: mensaje claro y datos conservados", "NO se ha confirmado" in page.inner_text("#form-status") and page.input_value("#full_name") == "Pedro Gómez Salas")
    page.locator("#panel").screenshot(path=os.path.join(OUT, "mobile-server-error.png"))
    page.click("#submit-btn")
    page.wait_for_selector("#confirm:not([hidden])", timeout=8000)
    check("reintento exitoso tras error", page.inner_text("#c-folio").startswith("LANDA-"))
    page.screenshot(path=os.path.join(OUT, "mobile-confirmation.png"))
    ctx.close()

    # ------------------------------------------------ páginas legales
    ctx = browser.new_context(viewport={"width": 1280, "height": 900}, reduced_motion="reduce")
    page = ctx.new_page()
    page.goto(BASE + "/aviso-privacidad.html", wait_until="networkidle")
    check("aviso de privacidad accesible", "Aviso de Privacidad" in page.inner_text("h1"))
    check("aviso marca datos pendientes", page.locator(".pending").count() > 0 and page.locator("#draft-banner").is_visible())
    page.screenshot(path=os.path.join(OUT, "privacy.png"), full_page=True)
    page.goto(BASE + "/derechos.html", wait_until="networkidle")
    page.fill("#r-folio", folio); page.fill("#r-email", "laura@example.com")
    page.click("#r-btn")
    page.wait_for_timeout(600)
    check("baja de publicidad responde", "preferencia" in page.inner_text("#r-status"))
    ctx.close()

    # ------------------------------------------------ panel administrativo
    ctx = browser.new_context(viewport={"width": 1440, "height": 900}, accept_downloads=True)
    page = ctx.new_page()
    page.goto(BASE + "/admin.html", wait_until="networkidle")
    check("admin: sin sesión muestra login", page.locator("#login-view").is_visible() and page.locator("#app-view").is_hidden())
    page.screenshot(path=os.path.join(OUT, "admin-login.png"))
    page.fill("#l-user", USER); page.fill("#l-pass", "contraseña-incorrecta-1")
    page.click("#login-btn"); page.wait_for_timeout(600)
    check("admin: contraseña incorrecta rechazada", "incorrectos" in page.inner_text("#login-status"))
    page.fill("#l-pass", PASS); page.click("#login-btn")
    page.wait_for_selector("#app-view:not([hidden])", timeout=8000)
    page.wait_for_timeout(800)
    check("admin: dashboard con total de registros", int(page.inner_text("#s-total")) >= 2)
    page.screenshot(path=os.path.join(OUT, "admin-dashboard.png"), full_page=True)
    page.fill("#f-q", "Laura"); page.click("#filters button[type=submit]"); page.wait_for_timeout(600)
    check("admin: búsqueda por nombre", "registro(s)" in page.inner_text("#result-note") and "Laura" in page.inner_text("#tbody"), page.inner_text("#result-note"))
    page.click("#tbody button"); page.wait_for_selector("#detail[open]")
    page.screenshot(path=os.path.join(OUT, "admin-detail.png"))
    page.select_option("#d-status", "confirmed"); page.click("#d-save"); page.wait_for_timeout(600)
    check("admin: cambio de estado", "Estado guardado" in page.inner_text("#d-status-msg"))
    page.click("#d-close")
    page.fill("#f-q", ""); page.click("#filters button[type=submit]"); page.wait_for_timeout(500)
    with page.expect_download() as dl:
        page.click("#btn-export")
    xlsx_path = os.path.join(OUT, "expo-landa-registros.xlsx")
    dl.value.save_as(xlsx_path)
    wb = load_workbook(xlsx_path)
    rows = list(wb["Registros"].iter_rows(values_only=True))
    check("admin: Excel descargado con registros reales", len(rows) >= 3 and rows[0][0] == "Folio", len(rows) - 1)
    check("admin: Excel contiene el folio registrado", any(r[0] == folio for r in rows[1:]))
    page.click("#tab-audit"); page.wait_for_timeout(500)
    page.screenshot(path=os.path.join(OUT, "admin-audit.png"))
    page.click("#btn-logout"); page.wait_for_timeout(600)
    check("admin: cierre de sesión", page.locator("#login-view").is_visible())
    page.goto(BASE + "/admin.html", wait_until="networkidle")
    check("admin: tras logout no hay acceso", page.locator("#login-view").is_visible())
    ctx.close()
    browser.close()

failed = [n for n, ok in results if not ok]
print("\n%d/%d verificaciones correctas" % (len(results) - len(failed), len(results)))
if failed:
    print("FALLARON:", failed)
    sys.exit(1)
