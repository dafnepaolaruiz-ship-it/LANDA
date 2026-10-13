"""Pruebas del panel administrativo, Excel y seguridad (puntos 15-30)."""
import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

import auth
import models
from config import MARKETING_TEXT_VERSION
from database import SessionLocal
from tests.conftest import PASSWORD, admin_client, create_user, login, register

ADMIN_GET = ["/api/admin/me", "/api/admin/stats", "/api/admin/audit", "/api/admin/registrations/LANDA-AAAA-BBBB"]
ADMIN_POST = [
    ("post", "/api/admin/registrations/search", {}),
    ("post", "/api/admin/export", {}),
    ("patch", "/api/admin/registrations/LANDA-AAAA-BBBB/status", {"status": "confirmed"}),
    ("post", "/api/admin/registrations/LANDA-AAAA-BBBB/anonymize", {"confirm_folio": "x"}),
    ("post", "/api/admin/registrations/LANDA-AAAA-BBBB/delete", {"confirm_folio": "x"}),
    ("post", "/api/admin/registrations/LANDA-AAAA-BBBB/revoke-marketing", None),
    ("post", "/api/admin/logout", None),
    ("post", "/api/admin/change-password", {"current_password": "a", "new_password": "b"}),
]


# ---- 15, 16, 27: sin sesión no hay acceso a nada
@pytest.mark.parametrize("path", ADMIN_GET)
def test_unauthenticated_get_blocked(client, path):
    assert client.get(path).status_code == 401


@pytest.mark.parametrize("method,path,body", ADMIN_POST)
def test_unauthenticated_write_and_export_blocked(client, method, path, body):
    r = getattr(client, method)(path, json=body) if body is not None else getattr(client, method)(path)
    assert r.status_code in (401, 403) and r.status_code != 200
    assert "xlsx" not in r.headers.get("content-type", "")


def test_forged_cookie_rejected(client):
    client.cookies.set("landa_session", "inventada")
    assert client.get("/api/admin/stats").status_code == 401


# ---- login, hashing, bloqueo
def test_login_wrong_password_generic_message(client):
    create_user()
    a = login(client, password="mala-contraseña-1")
    b = login(client, username="noexiste")
    assert a.status_code == b.status_code == 401 and a.json() == b.json()


def test_passwords_stored_with_argon2(client):
    create_user()
    with SessionLocal() as db:
        h = db.execute(select(models.AdminUser)).scalar_one().password_hash
    assert h.startswith("$argon2id$") and PASSWORD not in h


def test_account_lockout_after_repeated_failures(client):
    create_user()
    for _ in range(5):
        login(client, password="incorrecta-123456")
    r = login(client)  # incluso con la contraseña correcta
    assert r.status_code == 429


def test_login_rate_limit_per_ip(client):
    for _ in range(10):
        login(client, username="x", password="y")
    assert login(client, username="x", password="y").status_code == 429


def test_session_cookie_attributes(client):
    create_user()
    r = login(client)
    sc = r.headers["set-cookie"].lower()
    assert "httponly" in sc and "samesite=strict" in sc and "path=/" in sc
    assert "csrf_token" in r.json() and PASSWORD not in r.text


def test_secure_flag_set_when_enabled(client, monkeypatch):
    import config
    import routes.admin as ra

    monkeypatch.setattr(ra.settings, "cookie_secure", True)
    create_user()
    r = login(client)
    assert "secure" in r.headers["set-cookie"].lower() and r.headers["set-cookie"].startswith("__Host-")


# ---- 18: expiración
def test_session_expires_when_idle(client):
    ac = admin_client(client)
    assert ac.get("/api/admin/me").status_code == 200
    with SessionLocal() as db:
        s = db.execute(select(models.AdminSession)).scalar_one()
        s.last_seen_at = models.utcnow() - timedelta(minutes=auth.settings.session_idle_minutes + 1)
        db.commit()
    assert ac.get("/api/admin/me").status_code == 401


def test_session_absolute_expiry(client):
    ac = admin_client(client)
    with SessionLocal() as db:
        s = db.execute(select(models.AdminSession)).scalar_one()
        s.expires_at = models.utcnow() - timedelta(seconds=1)
        db.commit()
    assert ac.get("/api/admin/me").status_code == 401


def test_logout_invalidates_server_session(client):
    ac = admin_client(client)
    token = ac.cookies.get("landa_session")
    assert ac.post("/api/admin/logout").status_code == 200
    ac.cookies.set("landa_session", token)  # reutilizar el token anterior
    assert ac.get("/api/admin/me").status_code == 401


# ---- CSRF / origen
def test_csrf_required_for_writes(client):
    ac = admin_client(client)
    del ac.headers["X-CSRF-Token"]
    assert ac.post("/api/admin/registrations/search", json={}).status_code == 403
    ac.headers["X-CSRF-Token"] = "falso"
    assert ac.post("/api/admin/registrations/search", json={}).status_code == 403


def test_foreign_origin_rejected(client):
    ac = admin_client(client)
    r = ac.post("/api/admin/registrations/search", json={}, headers={"Origin": "https://evil.example"})
    assert r.status_code == 403


# ---- 17: roles
def test_viewer_cannot_modify_or_export(client):
    register(client)
    folio = register(client, full_name="Otra Persona Prueba", phone="449 777 8899").json()["folio"]
    vc = admin_client(client, role="viewer", username="lector")
    assert vc.get("/api/admin/stats").status_code == 200  # lectura permitida
    assert vc.post("/api/admin/registrations/search", json={}).status_code == 200
    assert vc.patch("/api/admin/registrations/%s/status" % folio, json={"status": "cancelled"}).status_code == 403
    assert vc.post("/api/admin/export", json={}).status_code == 403
    assert vc.post("/api/admin/registrations/%s/anonymize" % folio, json={"confirm_folio": folio}).status_code == 403
    with SessionLocal() as db:
        assert db.execute(select(models.Registration).where(models.Registration.folio == folio)).scalar_one().status == "pending"


# ---- 19: filtros y paginación
def _seed(client, n=12):
    for i in range(n):
        register(
            client,
            full_name="Asistente Número %02d" % i,
            phone="449 100 %04d" % (4000 + i),
            email="asistente%d@example.com" % i,
            attendees=1 + (i % 4),
        )


def test_pagination_and_filters(client):
    _seed(client, 12)
    ac = admin_client(client)
    r = ac.post("/api/admin/registrations/search", json={"page": 1, "page_size": 5}).json()
    assert r["total"] == 12 and len(r["items"]) == 5 and r["pages"] == 3
    last = ac.post("/api/admin/registrations/search", json={"page": 3, "page_size": 5}).json()
    assert len(last["items"]) == 2
    assert ac.post("/api/admin/registrations/search", json={"q": "Número 07"}).json()["total"] == 1
    assert ac.post("/api/admin/registrations/search", json={"q": "4003"}).json()["total"] == 1  # teléfono
    assert ac.post("/api/admin/registrations/search", json={"q": "asistente5@example"}).json()["total"] == 1  # correo
    assert ac.post("/api/admin/registrations/search", json={"min_attendees": 4}).json()["total"] == 3
    assert ac.post("/api/admin/registrations/search", json={"status": "confirmed"}).json()["total"] == 0
    assert ac.post("/api/admin/registrations/search", json={"date_from": "2000-01-01", "date_to": "2999-01-01"}).json()["total"] == 12
    assert ac.post("/api/admin/registrations/search", json={"date_from": "2999-01-01"}).json()["total"] == 0
    assert ac.post("/api/admin/registrations/search", json={"page_size": 5000}).status_code == 422


def test_search_wildcards_are_literal(client):
    _seed(client, 3)
    ac = admin_client(client)
    assert ac.post("/api/admin/registrations/search", json={"q": "%"}).json()["total"] == 0
    assert ac.post("/api/admin/registrations/search", json={"q": "' OR 1=1 --"}).json()["total"] == 0


# ---- 20: cambio de estado persiste + auditoría + criterio de asistentes
def test_status_change_persists_and_is_audited(client):
    folio = register(client, attendees=4).json()["folio"]
    register(client, attendees=2, full_name="Segunda Persona Prueba", phone="449 600 0000")
    ac = admin_client(client)
    r = ac.patch("/api/admin/registrations/%s/status" % folio, json={"status": "confirmed"})
    assert r.status_code == 200 and r.json()["status"] == "confirmed"
    with SessionLocal() as db:
        assert db.execute(select(models.Registration).where(models.Registration.folio == folio)).scalar_one().status == "confirmed"
        a = db.execute(select(models.AuditLog).where(models.AuditLog.action == "status_changed")).scalar_one()
        assert a.actor == "organizadora" and a.target_folio == folio and "maria" not in (a.detail or "").lower()
    ac.patch("/api/admin/registrations/%s/status" % folio, json={"status": "cancelled"})
    s = ac.get("/api/admin/stats").json()
    assert s["total_registrations"] == 2 and s["total_attendees"] == 2  # excluye cancelados
    assert s["by_status"]["cancelled"]["registrations"] == 1
    assert ac.patch("/api/admin/registrations/%s/status" % folio, json={"status": "hackeado"}).status_code == 422


# ---- 21-24: Excel
def _load(resp):
    return load_workbook(io.BytesIO(resp.content))


def test_export_contains_expected_data_and_all_filtered_rows(client):
    _seed(client, 30)  # más que una página
    ac = admin_client(client)
    r = ac.post("/api/admin/export", json={})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"] and "attachment" in r.headers["content-disposition"]
    ws = _load(r)["Registros"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Folio" and "Total de asistentes" in rows[0] and "Preferencia de publicidad" in rows[0]
    assert len(rows) - 1 == 30  # todos, no solo la página visible
    sample = [x for x in rows[1:] if x[2] == "Asistente Número 05"][0]
    assert sample[3] == "4491004005" and sample[4] == "asistente5@example.com" and sample[5] == 2 and sample[7] == "Pendiente"
    # filtro: solo los de 4 asistentes
    f = _load(ac.post("/api/admin/export", json={"min_attendees": 4}))["Registros"]
    assert len(list(f.iter_rows(values_only=True))) - 1 == len([i for i in range(30) if 1 + i % 4 >= 4])


def test_export_respects_status_filter_and_is_audited(client):
    f1 = register(client).json()["folio"]
    register(client, full_name="Otra Persona Prueba", phone="449 888 0000")
    ac = admin_client(client)
    ac.patch("/api/admin/registrations/%s/status" % f1, json={"status": "confirmed"})
    ws = _load(ac.post("/api/admin/export", json={"status": "confirmed"}))["Registros"]
    assert [r[0] for r in list(ws.iter_rows(values_only=True))[1:]] == [f1]
    with SessionLocal() as db:
        assert db.execute(select(models.AuditLog).where(models.AuditLog.action == "export_xlsx")).scalars().first()


@pytest.mark.parametrize("evil", ["=HYPERLINK(\"http://evil\",\"x\")", "+cmd|' /C calc'!A0", "-2+3", "@SUM(1+1)", "\t=1+1", "  =1+1"])
def test_export_neutralizes_formula_injection(client, evil):
    register(client, full_name=evil + " Persona")
    ac = admin_client(client)
    wb = _load(ac.post("/api/admin/export", json={}))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate, c.value)
                if isinstance(c.value, str):
                    assert not c.value.startswith(("=", "+", "-", "@", "\t", "\r")), c.value


def test_export_has_no_secrets_or_session_data(client):
    register(client, marketing_opt_in=True, marketing_text_version=MARKETING_TEXT_VERSION)
    ac = admin_client(client)
    token = ac.cookies.get("landa_session")
    csrf = ac.headers["X-CSRF-Token"]
    r = ac.post("/api/admin/export", json={})
    wb = _load(r)
    text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None).lower()
    with SessionLocal() as db:
        h = db.execute(select(models.AdminUser)).scalar_one().password_hash
    for forbidden in (token.lower(), csrf.lower(), h.lower(), "argon2", PASSWORD.lower(), "token", "hash", "secret", "idempotency"):
        assert forbidden not in text, forbidden
    headers = [c.value for c in wb["Registros"][1]]
    assert not any(k in " ".join(headers).lower() for k in ("hash", "token", "sesión", "password", "contraseña"))


# ---- 25, 26, 28: validación, XSS, PII en URL/errores
def test_stored_html_is_returned_as_plain_data_and_panel_uses_textcontent(client):
    payload = "<img src=x onerror=alert(1)> Persona"
    assert register(client, full_name=payload).status_code == 201
    ac = admin_client(client)
    item = ac.post("/api/admin/registrations/search", json={}).json()["items"][0]
    assert item["full_name"] == payload  # se almacena tal cual; el escape ocurre al renderizar
    assert "application/json" in ac.post("/api/admin/registrations/search", json={}).headers["content-type"]
    from pathlib import Path

    js = (Path(__file__).resolve().parents[2] / "frontend" / "admin.js").read_text(encoding="utf-8")
    assert "innerHTML" not in js, "admin.js no debe usar innerHTML con datos de usuarios"
    assert "insertAdjacentHTML" not in js and "document.write" not in js


def test_no_personal_data_in_urls_or_error_bodies(client):
    r = register(client, email="secreto-persona@example.com", phone="abc")
    assert r.status_code == 422
    assert "secreto-persona" not in r.text
    assert r.url.query == b""
    assert client.get("/api/registrations").status_code in (404, 405)


def test_error_responses_hide_internals(client):
    ac = admin_client(client)
    r = ac.get("/api/admin/registrations/no-es-folio")
    assert r.status_code == 404 and "Traceback" not in r.text and "sqlalchemy" not in r.text.lower()


def test_oversized_request_rejected(client):
    r = client.post("/api/registrations", content=b"{" + b" " * 40000 + b"}", headers={"Content-Type": "application/json"})
    assert r.status_code == 413


def test_sql_injection_attempt_is_inert(client):
    r = register(client, full_name="Robert'); DROP TABLE registrations;-- Persona")
    assert r.status_code == 201
    with SessionLocal() as db:
        assert db.execute(select(models.Registration)).scalars().first().full_name.startswith("Robert")


def test_security_headers_present(client):
    r = client.get("/api/config")
    for h in ("content-security-policy", "x-content-type-options", "referrer-policy", "permissions-policy"):
        assert h in r.headers
    assert "frame-ancestors 'none'" in r.headers["content-security-policy"]


def test_admin_responses_are_not_cacheable(client):
    ac = admin_client(client)
    assert ac.get("/api/admin/stats").headers["cache-control"] == "no-store"


def test_cors_limited_to_allowed_origins(client):
    ok = client.options("/api/registrations", headers={"Origin": "http://localhost:8000", "Access-Control-Request-Method": "POST"})
    bad = client.options("/api/registrations", headers={"Origin": "https://evil.example", "Access-Control-Request-Method": "POST"})
    assert ok.headers.get("access-control-allow-origin") == "http://localhost:8000"
    assert "access-control-allow-origin" not in bad.headers


def test_docs_disabled_in_production_setting():
    from main import app

    # En desarrollo existen; en producción se desactivan (ver main.py)
    import inspect, main

    assert "None if settings.is_production" in inspect.getsource(main)


# ---- 29: preferencias publicitarias
def test_marketing_preferences_respected_in_filters_and_export(client):
    a = register(client, marketing_opt_in=True, marketing_text_version=MARKETING_TEXT_VERSION).json()["folio"]
    b = register(client, full_name="Sin Publicidad Persona", phone="449 999 1111", email="b@example.com").json()["folio"]
    ac = admin_client(client)
    allowed = ac.post("/api/admin/registrations/search", json={"marketing": "allowed"}).json()
    assert [i["folio"] for i in allowed["items"]] == [a]
    assert ac.post("/api/admin/registrations/%s/revoke-marketing" % a).json()["marketing_allowed"] is False
    assert ac.post("/api/admin/registrations/search", json={"marketing": "allowed"}).json()["total"] == 0
    ws = _load(ac.post("/api/admin/export", json={}))["Registros"]
    col = [c.value for c in ws[1]].index("Preferencia de publicidad")
    assert {r[col] for r in ws.iter_rows(min_row=2, values_only=True)} == {"No"}
    d = ac.get("/api/admin/registrations/%s" % a).json()
    assert d["marketing_opt_in"] is True and d["marketing_revoked_at"] and d["marketing_consent_at"]


# ---- Conservación / ARCO
def test_anonymize_requires_confirmation_and_removes_pii(client):
    folio = register(client, attendees=4).json()["folio"]
    ac = admin_client(client)
    assert ac.post("/api/admin/registrations/%s/anonymize" % folio, json={"confirm_folio": "LANDA-XXXX-XXXX"}).status_code == 400
    r = ac.post("/api/admin/registrations/%s/anonymize" % folio, json={"confirm_folio": folio, "reason": "Solicitud ARCO"})
    assert r.status_code == 200
    d = r.json()
    assert d["full_name"] is None and d["phone"] is None and d["email"] is None and d["anonymized"] and d["attendees"] == 4
    assert ac.get("/api/admin/stats").json()["total_attendees"] == 4  # conserva el dato estadístico


def test_delete_registration_and_audit(client):
    folio = register(client).json()["folio"]
    ac = admin_client(client)
    assert ac.post("/api/admin/registrations/%s/delete" % folio, json={"confirm_folio": folio}).status_code == 200
    assert ac.get("/api/admin/registrations/%s" % folio).status_code == 404
    with SessionLocal() as db:
        a = db.execute(select(models.AuditLog).where(models.AuditLog.action == "deleted")).scalar_one()
        assert a.target_folio == folio


def test_cli_anonymize_expired(client, monkeypatch):
    import cli
    from argparse import Namespace

    folio = register(client).json()["folio"]
    register(client, full_name="Registro Reciente Prueba", phone="449 121 2121")
    with SessionLocal() as db:
        r = db.execute(select(models.Registration).where(models.Registration.folio == folio)).scalar_one()
        r.created_at = models.utcnow() - timedelta(days=800)
        db.commit()
    assert cli.cmd_anonymize(Namespace(months=0, dry_run=False)) == 1  # exige periodo explícito
    cli.cmd_anonymize(Namespace(months=12, dry_run=True))
    with SessionLocal() as db:
        assert db.execute(select(models.Registration).where(models.Registration.anonymized_at.is_not(None))).first() is None
    cli.cmd_anonymize(Namespace(months=12, dry_run=False))
    with SessionLocal() as db:
        rows = db.execute(select(models.Registration)).scalars().all()
        assert sum(1 for r in rows if r.anonymized_at) == 1 and sum(1 for r in rows if r.email) == 1


def test_change_password_rotates_sessions(client):
    ac = admin_client(client)
    old = ac.cookies.get("landa_session")
    assert ac.post("/api/admin/change-password", json={"current_password": PASSWORD, "new_password": "corta"}).status_code == 422
    assert ac.post("/api/admin/change-password", json={"current_password": "equivocada-123", "new_password": "Nueva-Contraseña-456"}).status_code == 400
    r = ac.post("/api/admin/change-password", json={"current_password": PASSWORD, "new_password": "Nueva-Contraseña-456"})
    assert r.status_code == 200
    ac.headers["X-CSRF-Token"] = r.json()["csrf_token"]
    assert ac.get("/api/admin/me").status_code == 200
    ac.cookies.clear()
    ac.cookies.set("landa_session", old)
    assert ac.get("/api/admin/me").status_code == 401  # la sesión anterior ya no sirve


def test_public_config_has_no_secrets(client):
    t = client.get("/api/config").text.lower()
    for bad in ("secret", "password", "database", "turnstile_secret"):
        assert bad not in t


def test_privacy_placeholders_when_not_configured(client):
    from config import Settings, get_settings
    from main import app
    dummy = Settings(
        form_signing_secret="x" * 48,
        privacy_controller_name="",
        privacy_controller_address="",
        privacy_contact_email="",
        privacy_retention_text="",
    )
    app.dependency_overrides[get_settings] = lambda: dummy
    try:
        c = client.get("/api/config").json()
        assert c["privacy"]["legal_configured"] is False and "POR CONFIGURAR" in c["privacy"]["controller_name"]
    finally:
        app.dependency_overrides.pop(get_settings, None)


def test_privacy_configured_when_settings_provided(client):
    from config import Settings, get_settings
    from main import app
    dummy = Settings(
        form_signing_secret="x" * 48,
        privacy_controller_name="Landa Entorno Residencial",
        privacy_controller_address="Blvd. El Chichimeco s/n",
        privacy_contact_email="contacto@landa.mx",
        privacy_retention_text="12 meses",
    )
    app.dependency_overrides[get_settings] = lambda: dummy
    try:
        c = client.get("/api/config").json()
        assert c["privacy"]["legal_configured"] is True
        assert c["privacy"]["controller_name"] == "Landa Entorno Residencial"
    finally:
        app.dependency_overrides.pop(get_settings, None)
